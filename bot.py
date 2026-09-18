from telegram import Update
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, filters, ConversationHandler, ContextTypes
from docx import Document
from datetime import date, datetime, timedelta
from dateutil.relativedelta import relativedelta
import os

TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "8612996676:AAH4cviwntK333WC3cuTp3d9Jvbhf60nDf0")
TEMPLATES_DIR = "templates"
OUTPUT_DIR = "output"

(CHOOSE_TYPE, CHOOSE_CITIZENSHIP, CHOOSE_STUDENT, CHOOSE_POSITION,
 ASK_ZANYATOST, ASK_STAVKA,
 ASK_FIO, ASK_BIRTH, ASK_PASSPORT_SERIES, ASK_PASSPORT_NUMBER, ASK_PASSPORT_ISSUED,
 ASK_ADDRESS_REG, ASK_ADDRESS_LIVE, ASK_SALARY, ASK_START_DATE, ASK_CONTRACT_TERM,
 ASK_UNIVERSITY, ASK_STUDY_END) = range(18)

OBYAZANNOSTI = {
    "директор": "— Руководство деятельностью Общества.\n— Подписание договоров и финансовых документов.\n— Утверждение штатного расписания.\n— Издание приказов и распоряжений.",
    "заместитель директора по экономике и финансам": "— Сведение и анализ инвентаризаций.\n— Работа с поставщиками.\n— Бюджетирование и управление себестоимостью.\n— Подготовка финансовой отчётности.",
    "заместитель директора по производству": "— Хозяйственная часть.\n— Поиск и подбор персонала.\n— Координация работы отделов, водителей, управляющих.\n— Контроль стандартов.",
    "экономист по маркетингу и кадровой работе": "— Кадровое делопроизводство.\n— Маркетинг (продвижение, соцсети, акции).\n— Сбор данных для бухгалтерии.",
    "специалист по стандартам и обучению": "— Разработка стандартов работы.\n— Обучение и аттестация персонала.\n— Контроль соблюдения стандартов.",
    "управляющий отдела": "— Приготовление блюд, контроль качества.\n— Управление отделом.\n— Заказ сырья, работа с поставщиками.\n— Инвентаризации, отчётность.",
    "повар-наставник": "— Обучение стажёров и сотрудников.\n— Контроль соблюдения техкарт.\n— Проведение аттестаций.\n— Участие в инвентаризациях.",
    "старший повар": "— Контроль качества блюд.\n— Расстановка поваров.\n— Формирование заявок на сырьё.\n— Участие в инвентаризациях.",
    "повар 3 разряда": "— Приготовление блюд согласно техкартам.\n— Соблюдение санитарных норм.\n— Поддержание чистоты.\n— Участие в инвентаризациях.",
    "старший кассир": "— Контроль работы кассиров.\n— Инкассация, сверка касс.\n— Сервировка, сбор заказов, доготовка.\n— Передача заказов водителям и курьерам.\n— Участие в инвентаризациях.",
    "кассир-администратор": "— Открытие/закрытие смены.\n— Координация персонала зала.\n— Работа на кассе.\n— Сервировка, сбор заказов, доготовка.\n— Передача заказов.\n— Участие в инвентаризациях.",
    "кассир": "— Работа на кассе.\n— Сервировка и сбор заказов.\n— Доготовка блюд.\n— Передача заказов водителям и курьерам.\n— Участие в инвентаризациях.",
    "водитель автомобиля": "— Доставка продукции и сырья.\n— Контроль технического состояния авто.\n— Оформление путевых листов.",
    "кухонный рабочий": "— Мойка посуды, инвентаря.\n— Уборка кухни и зала.\n— Транспортировка продуктов и отходов.",
}

MAT_OTVETSTVENNOST = {
    "директор": "6.2. Работник несет полную материальную ответственность за имущество Общества.",
    "управляющий отдела": "6.2. Работник несет полную индивидуальную материальную ответственность за имущество отдела, денежные средства и ТМЦ.",
    "повар-наставник": "6.2. Работник несет полную индивидуальную материальную ответственность за вверенные ТМЦ.",
    "старший повар": "6.2. Работник несет полную индивидуальную материальную ответственность за вверенные ТМЦ.",
    "повар 3 разряда": "6.2. Работник несет полную индивидуальную материальную ответственность за вверенные ТМЦ.",
    "старший кассир": "6.2. Работник несет полную индивидуальную материальную ответственность за денежные средства и ТМЦ.",
    "кассир-администратор": "6.2. Работник несет полную индивидуальную материальную ответственность за денежные средства и ТМЦ.",
    "кассир": "6.2. Работник несет полную индивидуальную материальную ответственность за денежные средства и ТМЦ.",
    "водитель автомобиля": "6.2. Работник несет полную индивидуальную материальную ответственность за автомобиль и груз.",
}

PERECHEN_115 = ["повар 3 разряда", "старший повар", "повар-наставник", "водитель автомобиля"]
DOLZHNOSTI_LIST = "\n".join([f"• {d}" for d in OBYAZANNOSTI.keys()])

def get_template_name(context):
    cit = context.user_data.get('citizenship', 'РБ')
    typ = context.user_data.get('type', 'bessrochny')
    student = context.user_data.get('student', False)
    pos = context.user_data.get('position', '')
    if student:
        return "template_student.docx"
    elif cit in ['РБ', 'ЕАЭС']:
        return "template_rb_contract.docx" if typ == 'contract' else "template_rb_bessrochny.docx"
    elif pos in PERECHEN_115:
        return "template_perechen115.docx"
    else:
        return "template_specrazreshenie.docx"

def get_end_date(context):
    cit = context.user_data.get('citizenship', 'РБ')
    pos = context.user_data.get('position', '')
    if cit not in ['РБ', 'ЕАЭС'] and pos in PERECHEN_115:
        return f"31.12.{date.today().year}"
    return context.user_data.get('end_date', '')

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("👋 Выберите тип договора:\n1️⃣ Бессрочный\n2️⃣ Контракт (срочный)")
    return CHOOSE_TYPE

async def choose_type(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().lower()
    if text in ['1', 'бессрочный']:
        context.user_data['type'] = 'bessrochny'
    elif text in ['2', 'контракт', 'срочный']:
        context.user_data['type'] = 'contract'
    else:
        await update.message.reply_text("Пожалуйста, выберите: 1 или 2")
        return CHOOSE_TYPE
    await update.message.reply_text("🌍 Гражданство:\n1️⃣ РБ\n2️⃣ РФ\n3️⃣ Казахстан\n4️⃣ Армения\n5️⃣ Киргизия\n6️⃣ Иное")
    return CHOOSE_CITIZENSHIP

async def choose_citizenship(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().lower()
    eaeu = ['2', '3', '4', '5', 'рф', 'россия', 'казахстан', 'армения', 'киргизия']
    rb = ['1', 'рб', 'беларусь', 'белорус']
    if any(c in text for c in rb):
        context.user_data['citizenship'] = 'РБ'
    elif any(c in text for c in eaeu):
        context.user_data['citizenship'] = 'ЕАЭС'
    else:
        context.user_data['citizenship'] = text
        await update.message.reply_text("🎓 Студент дневной формы в РБ? (да/нет)")
        return CHOOSE_STUDENT
    await update.message.reply_text(f"👤 Должность:\n{DOLZHNOSTI_LIST}")
    return CHOOSE_POSITION

async def choose_student(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message.text.strip().lower() in ['да', 'yes', '1']:
        context.user_data['student'] = True
        await update.message.reply_text("🏫 Название учреждения образования:")
        return ASK_UNIVERSITY
    context.user_data['student'] = False
    await update.message.reply_text(f"👤 Должность:\n{DOLZHNOSTI_LIST}")
    return CHOOSE_POSITION

async def ask_university(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['university'] = update.message.text.strip()
    await update.message.reply_text("📅 Дата окончания обучения (ДД.ММ.ГГГГ):")
    return ASK_STUDY_END

async def ask_study_end(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['study_end'] = update.message.text.strip()
    await update.message.reply_text(f"👤 Должность:\n{DOLZHNOSTI_LIST}")
    return CHOOSE_POSITION

async def choose_position(update: Update, context: ContextTypes.DEFAULT_TYPE):
    pos = update.message.text.strip().lower()
    if pos in OBYAZANNOSTI:
        context.user_data['position'] = pos
        await update.message.reply_text("💼 Занятость:\n1️⃣ Основное место\n2️⃣ Внешнее совместительство\n3️⃣ Внутреннее совместительство")
        return ASK_ZANYATOST
    await update.message.reply_text(f"Выберите из списка:\n{DOLZHNOSTI_LIST}")
    return CHOOSE_POSITION

async def ask_zanyatost(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().lower()
    opts = {'1': 'по основному месту работы', '2': 'по внешнему совместительству', '3': 'по внутреннему совместительству',
            'основное': 'по основному месту работы', 'внешнее': 'по внешнему совместительству', 'внутреннее': 'по внутреннему совместительству'}
    if text in opts:
        context.user_data['zanyatost'] = opts[text]
        await update.message.reply_text("📊 Ставка:\n1️⃣ 1.0\n2️⃣ 0.75\n3️⃣ 0.5\n4️⃣ 0.25")
        return ASK_STAVKA
    await update.message.reply_text("1. Основное\n2. Внешнее\n3. Внутреннее")
    return ASK_ZANYATOST

async def ask_stavka(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    stavki = {'1': '1.0 (полная)', '2': '0.75', '3': '0.5', '4': '0.25'}
    if text in stavki:
        context.user_data['stavka'] = stavki[text]
        await update.message.reply_text("👤 ФИО полностью:")
        return ASK_FIO
    await update.message.reply_text("1, 2, 3 или 4")
    return ASK_STAVKA

async def ask_fio(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['fio'] = update.message.text.strip()
    await update.message.reply_text("🎂 Дата рождения (ДД.ММ.ГГГГ):")
    return ASK_BIRTH

async def ask_birth(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['birth'] = update.message.text.strip()
    await update.message.reply_text("📘 Серия паспорта:")
    return ASK_PASSPORT_SERIES

async def ask_passport_series(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['pass_ser'] = update.message.text.strip().upper()
    await update.message.reply_text("📘 Номер паспорта:")
    return ASK_PASSPORT_NUMBER

async def ask_passport_number(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['pass_num'] = update.message.text.strip()
    await update.message.reply_text("🏛 Кем выдан:")
    return ASK_PASSPORT_ISSUED

async def ask_passport_issued(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['pass_issued'] = update.message.text.strip()
    await update.message.reply_text("📍 Адрес регистрации:")
    return ASK_ADDRESS_REG

async def ask_address_reg(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['addr_reg'] = update.message.text.strip()
    await update.message.reply_text("🏠 Адрес проживания (или 'тот же'):")
    return ASK_ADDRESS_LIVE

async def ask_address_live(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ans = update.message.text.strip()
    context.user_data['addr_live'] = context.user_data['addr_reg'] if ans.lower() in ['тот же', 'тотже'] else ans
    await update.message.reply_text("💰 Оклад (BYN в месяц):")
    return ASK_SALARY

async def ask_salary(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['salary'] = update.message.text.strip()
    await update.message.reply_text("📅 Дата начала работы (ДД.ММ.ГГГГ):")
    return ASK_START_DATE

async def ask_start_date(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['start_date'] = update.message.text.strip()
    cit = context.user_data.get('citizenship', 'РБ')
    pos = context.user_data.get('position', '')
    if cit not in ['РБ', 'ЕАЭС'] and pos in PERECHEN_115:
        context.user_data['end_date'] = f"31.12.{date.today().year}"
        return await generate_doc(update, context)
    elif context.user_data.get('type') == 'contract':
        await update.message.reply_text("📅 Срок контракта в годах (1, 2, 3, 5 или 1.5):")
        return ASK_CONTRACT_TERM
    return await generate_doc(update, context)

async def ask_contract_term(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip().replace(',', '.')
    start_str = context.user_data.get('start_date', '')
    try:
        start_date = datetime.strptime(start_str, '%d.%m.%Y')
    except:
        await update.message.reply_text("❌ Ошибка в дате. /start заново.")
        return ConversationHandler.END
    try:
        years = float(text)
    except:
        await update.message.reply_text("Введите число (1, 2, 3, 5 или 1.5)")
        return ASK_CONTRACT_TERM
    end_date = start_date + relativedelta(years=int(years), months=int((years % 1) * 12)) - timedelta(days=1)
    context.user_data['end_date'] = end_date.strftime('%d.%m.%Y')
    return await generate_doc(update, context)

async def generate_doc(update: Update, context: ContextTypes.DEFAULT_TYPE):
    u = context.user_data
    template_name = get_template_name(context)
    template_path = os.path.join(TEMPLATES_DIR, template_name)

    if not os.path.exists(template_path):
        await update.message.reply_text(f"❌ Шаблон не найден: {template_name}")
        return ConversationHandler.END

    doc = Document(template_path)
    full = u.get('fio', '').split()
    fam = full[0] if len(full) > 0 else ''
    imya = full[1] if len(full) > 1 else ''
    otch = full[2] if len(full) > 2 else ''
    fio_imen = f"{imya} {otch}".strip()

    end_date = get_end_date(context)
    mat_block = MAT_OTVETSTVENNOST.get(u.get('position', ''), '')

    replacements = {
        '{{gorod}}': 'г. Витебск',
        '{{data_podpisania}}': '«___» _________ 202_ г.',
        '{{grazhdanstvo}}': 'гражданка' if fam.endswith('а') else 'гражданин',
        '{{FIO_rabotnika}}': u.get('fio', ''),
        '{{data_rozhdenia}}': u.get('birth', ''),
        '{{seria_pasporta}}': u.get('pass_ser', ''),
        '{{nomer_pasporta}}': u.get('pass_num', ''),
        '{{kem_vydan}}': u.get('pass_issued', ''),
        '{{adres_propiski}}': u.get('addr_reg', ''),
        '{{adres_prozhivania}}': u.get('addr_live', ''),
        '{{dolzhnost}}': u.get('position', ''),
        '{{mesto_raboty}}': 'ООО «ВИТАРМИКО»',
        '{{data_nachala}}': u.get('start_date', ''),
        '{{ispytatelny_srok}}': 'не устанавливается',
        '{{obyazannosti}}': OBYAZANNOSTI.get(u.get('position', ''), ''),
        '{{oklad}}': u.get('salary', ''),
        '{{blok_mat_otvetstvennost}}': mat_block,
        '{{FIO_rabotnika_imen}}': fio_imen,
        '{{FIO_rabotnika_fam}}': fam,
        '{{data_okonchania}}': end_date,
        '{{srok_deistvia}}': 'на определенный срок' if u.get('type') == 'contract' else 'на неопределенный срок',
        '{{university}}': u.get('university', ''),
        '{{study_end}}': u.get('study_end', ''),
        '{{tip_zanyatosti}}': u.get('zanyatost', 'по основному месту работы'),
        '{{stavka}}': u.get('stavka', '1.0 (полная)'),
    }

    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                p.text = p.text.replace(key, val)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for key, val in replacements.items():
                        if key in p.text:
                            p.text = p.text.replace(key, val)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filepath = os.path.join(OUTPUT_DIR, f"dogovor_{fam}.docx")
    doc.save(filepath)
    await update.message.reply_document(document=open(filepath, 'rb'), filename=f"Договор_{fam}.docx")
    await update.message.reply_text("✅ Готово! /start — новый договор.")
    return ConversationHandler.END

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("Отменено. /start для начала.")
    return ConversationHandler.END

if __name__ == '__main__':
    app = ApplicationBuilder().token(TOKEN).build()
    conv = ConversationHandler(
        entry_points=[CommandHandler('start', start)],
        states={
            CHOOSE_TYPE: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_type)],
            CHOOSE_CITIZENSHIP: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_citizenship)],
            CHOOSE_STUDENT: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_student)],
            ASK_UNIVERSITY: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_university)],
            ASK_STUDY_END: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_study_end)],
            CHOOSE_POSITION: [MessageHandler(filters.TEXT & ~filters.COMMAND, choose_position)],
            ASK_ZANYATOST: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_zanyatost)],
            ASK_STAVKA: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_stavka)],
            ASK_FIO: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_fio)],
            ASK_BIRTH: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_birth)],
            ASK_PASSPORT_SERIES: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_passport_series)],
            ASK_PASSPORT_NUMBER: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_passport_number)],
            ASK_PASSPORT_ISSUED: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_passport_issued)],
            ASK_ADDRESS_REG: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_address_reg)],
            ASK_ADDRESS_LIVE: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_address_live)],
            ASK_SALARY: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_salary)],
            ASK_START_DATE: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_start_date)],
            ASK_CONTRACT_TERM: [MessageHandler(filters.TEXT & ~filters.COMMAND, ask_contract_term)],
        },
        fallbacks=[CommandHandler('cancel', cancel)],
    )
    app.add_handler(conv)
    print("✅ Бот запущен!")
    app.run_polling()